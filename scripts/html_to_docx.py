"""Convert the assignment report HTML into an editable DOCX with embedded evidence."""

from __future__ import annotations

import sys
from html.parser import HTMLParser
from pathlib import Path
from urllib.parse import unquote, urlparse

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


NAVY = "18385A"
TEAL = "318C88"
GRAY = "59697A"


def set_cell_shading(cell, fill: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def set_cell_text(cell, text: str, header: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(8.5)
    if header:
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_shading(cell, NAVY)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


class ReportParser(HTMLParser):
    def __init__(self, document: Document):
        super().__init__(convert_charrefs=True)
        self.document = document
        self.paragraph = None
        self.paragraph_class = ""
        self.run_flags: list[str] = []
        self.in_cover = False
        self.in_table = False
        self.rows: list[list[str]] = []
        self.row: list[str] = []
        self.cell_text = ""
        self.cell_open = False
        self.list_stack: list[str] = []
        self.figure_count = 0

    def start_paragraph(self, style: str = "Normal", class_name: str = "") -> None:
        self.finish_paragraph()
        self.paragraph_class = class_name
        if style == "List Bullet":
            self.paragraph = self.document.add_paragraph(style="List Bullet")
        else:
            self.paragraph = self.document.add_paragraph(style=style)
        if self.in_cover or class_name in {"eyebrow", "subtitle", "author", "meta"}:
            self.paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if class_name == "author":
            self.paragraph.paragraph_format.space_before = Pt(18)
            self.paragraph.paragraph_format.space_after = Pt(3)
        if class_name == "eyebrow":
            self.paragraph.paragraph_format.space_after = Pt(18)

    def finish_paragraph(self) -> None:
        if self.paragraph is not None:
            self.paragraph = None
            self.paragraph_class = ""
            self.run_flags.clear()

    def add_text(self, text: str) -> None:
        if self.in_table:
            if self.cell_open:
                self.cell_text += text
            return
        if self.paragraph is None or not text:
            return
        run = self.paragraph.add_run(text)
        run.font.name = "Arial"
        if "strong" in self.run_flags or self.paragraph_class == "author":
            run.bold = True
        if "em" in self.run_flags:
            run.italic = True
        if "code" in self.run_flags:
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(55, 73, 91)
        if self.paragraph_class == "eyebrow":
            run.bold = True
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor.from_string(TEAL)
        elif self.paragraph_class == "author":
            run.font.size = Pt(16)
            run.font.color.rgb = RGBColor.from_string(NAVY)
        elif self.paragraph_class in {"meta", "small"}:
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor.from_string(GRAY)
        elif self.in_cover and self.paragraph.style.name == "Title":
            run.font.size = Pt(29)
            run.font.color.rgb = RGBColor.from_string(NAVY if "MODUL" not in text else TEAL)

    def handle_starttag(self, tag: str, attrs) -> None:
        attributes = dict(attrs)
        class_name = attributes.get("class", "")

        if tag == "section":
            if class_name == "cover":
                self.in_cover = True
            elif class_name == "evidence-page" and self.document.paragraphs:
                self.finish_paragraph()
                self.document.add_page_break()
            return
        if tag == "div":
            if class_name in {"eyebrow", "callout"}:
                self.start_paragraph("Normal", class_name)
            return
        if tag == "h1":
            self.start_paragraph("Title" if self.in_cover else "Heading 1")
            return
        if tag == "h2":
            self.start_paragraph("Subtitle" if self.in_cover else "Heading 2")
            return
        if tag == "h3":
            self.start_paragraph("Heading 3")
            return
        if tag == "p":
            style = "Caption" if class_name == "small" else "Normal"
            self.start_paragraph(style, class_name)
            return
        if tag in {"ul", "ol"}:
            self.list_stack.append(tag)
            return
        if tag == "li":
            self.start_paragraph("List Bullet")
            return
        if tag == "table":
            self.finish_paragraph()
            self.in_table = True
            self.rows = []
            return
        if tag == "tr":
            self.row = []
            return
        if tag in {"td", "th"}:
            self.cell_text = ""
            self.cell_open = True
            return
        if tag == "figure":
            self.figure_count += 1
            return
        if tag == "img":
            source = attributes.get("src", "")
            parsed = urlparse(source)
            image_path = Path(unquote(parsed.path.lstrip("/")))
            if image_path.exists():
                self.finish_paragraph()
                if self.figure_count > 1:
                    self.document.add_page_break()
                paragraph = self.document.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.space_before = Pt(2)
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.add_run().add_picture(str(image_path), width=Inches(5.75))
            return
        if tag == "figcaption":
            self.start_paragraph("Caption")
            return
        if tag in {"strong", "b", "code", "em", "i"}:
            self.run_flags.append("strong" if tag in {"strong", "b"} else "code" if tag == "code" else "em")
            return
        if tag == "br" and self.paragraph is not None:
            self.paragraph.add_run().add_break()

    def handle_endtag(self, tag: str) -> None:
        if tag in {"strong", "b", "code", "em", "i"}:
            flag = "strong" if tag in {"strong", "b"} else "code" if tag == "code" else "em"
            if flag in self.run_flags:
                self.run_flags.remove(flag)
            return
        if tag in {"p", "h1", "h2", "h3", "li", "figcaption"}:
            self.finish_paragraph()
            return
        if tag in {"td", "th"}:
            self.cell_open = False
            self.row.append(self.cell_text.strip())
            return
        if tag == "tr":
            if self.row:
                self.rows.append(self.row)
            return
        if tag == "table":
            self.in_table = False
            if self.rows:
                columns = max(len(row) for row in self.rows)
                table = self.document.add_table(rows=len(self.rows), cols=columns)
                table.style = "Table Grid"
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                for row_index, values in enumerate(self.rows):
                    for column_index in range(columns):
                        value = values[column_index] if column_index < len(values) else ""
                        set_cell_text(table.cell(row_index, column_index), value, row_index == 0)
            self.document.add_paragraph().paragraph_format.space_after = Pt(0)
            return
        if tag == "section" and self.in_cover:
            self.in_cover = False
            self.finish_paragraph()
            self.document.add_page_break()

    def handle_data(self, data: str) -> None:
        self.add_text(data)


def add_page_field(paragraph) -> None:
    run = paragraph.add_run("Herlina Dwi Septiana  |  Modul 6  |  Halaman ")
    run.font.name = "Arial"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(GRAY)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "1"
    run.append(text)
    field.append(run)
    paragraph._p.append(field)


def build(html_path: Path, output_path: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.58)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.66)
    section.right_margin = Inches(0.66)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    for style_name, size, color in (("Heading 1", 19, NAVY), ("Heading 2", 14, NAVY), ("Heading 3", 11, TEAL)):
        style = document.styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(5)

    title = document.styles["Title"]
    title.font.name = "Arial"
    title.font.size = Pt(29)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(NAVY)
    title.paragraph_format.space_before = Pt(60)
    title.paragraph_format.space_after = Pt(18)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = document.styles["Subtitle"]
    subtitle.font.name = "Arial"
    subtitle.font.size = Pt(18)
    subtitle.font.bold = True
    subtitle.font.color.rgb = RGBColor.from_string(TEAL)
    subtitle.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    caption = document.styles["Caption"]
    caption.font.name = "Arial"
    caption.font.size = Pt(8.5)
    caption.font.color.rgb = RGBColor.from_string(GRAY)
    caption.paragraph_format.space_after = Pt(5)

    header = section.header.paragraphs[0]
    header.text = "PRAKTIKUM PEMROGRAMAN 2  |  LAPORAN MODUL 6"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.name = "Arial"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(GRAY)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_field(footer)

    document.core_properties.author = "Herlina Dwi Septiana"
    document.core_properties.title = "Laporan Praktikum Modul 6 - Database 3"
    document.core_properties.subject = "Eloquent Relationship, Pivot Table, Factory, dan Seeder"

    parser = ReportParser(document)
    parser.feed(html_path.read_text(encoding="utf-8-sig"))
    parser.finish_paragraph()
    document.save(output_path)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("Usage: html_to_docx.py input.html output.docx")
    build(Path(sys.argv[1]), Path(sys.argv[2]))
